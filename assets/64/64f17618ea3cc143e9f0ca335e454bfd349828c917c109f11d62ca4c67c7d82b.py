from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
import re,ast,textwrap
p=Path('Thesis_Chapters1-5_Visuals_Code.docx');d=Document(p);base=Path('reports/opencv_illustrations')
pages=[re.sub(r'\s+','',x.extract_text()) for x in PdfReader(base/'render1'/f'{p.stem}.pdf').pages]
for x in d.paragraphs:
 if x.style.name.lower().startswith('toc'):
  ts=x._p.findall('.//'+qn('w:t'))
  if len(ts)<2:continue
  title=re.sub(r'\s+','',''.join(t.text or '' for t in ts[:-1]));matches=[i+1 for i,t in enumerate(pages) if title in t]
  assert matches,title
  ts[-1].text=str(matches[-1])
# Check that the reformatted method listings preserve the archived AST.
for file,func in [('bounded_state_local_16s.py','tick'),('v2_local_continuous_16s.py','update')]:
 source=ast.parse((base/file).read_text());original=next(n for n in ast.walk(source) if isinstance(n,ast.FunctionDef) and n.name==func)
 shown=next(x.text for x in d.paragraphs if x.text.startswith('def '+func+'('))
 assert ast.dump(original)==ast.dump(ast.parse(shown).body[0]),func
# Existing experiment-result text is preserved.
old=Document('Thesis_Chapters1-5_Methods_Expanded.docx')
def chapter5(doc):
 ps=doc.paragraphs;s=next(i for i,x in enumerate(ps) if x.style.name=='Heading 1' and x.text.startswith('5.'));e=next(i for i in range(s+1,len(ps)) if ps[i].style.name=='Heading 1');return [x.text for x in ps[s:e]]
assert chapter5(d)==chapter5(old)
assert [t._tbl.xml for t in d.tables]==[t._tbl.xml for t in old.tables]
d.save(p);print('TOC refreshed; listing ASTs and unchanged result tables verified.')
