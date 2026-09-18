from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
root=Path.cwd();base=root/'reports/opencv_illustrations';d=Document(root/'Thesis_Chapters1-5_Methods_Expanded.docx')
out=root/'Thesis_Chapters1-5_Visuals_Code.docx'
body=next(p for p in d.paragraphs if p.text.startswith('The final OpenCV controller estimates'))
head=next(p for p in d.paragraphs if p.style.name=='Heading 3' and p.text.startswith('4.4.1'))
def add(anchor,text,role='body'):
 template=head if role=='head' else body
 e=OxmlElement('w:p');e.append(deepcopy(template._p.pPr));anchor._p.addprevious(e);p=Paragraph(e,anchor._parent)
 r=p.add_run(text)
 if template.runs[0]._r.rPr is not None:r._r.insert(0,deepcopy(template.runs[0]._r.rPr))
 if role in ('caption','code'):
  p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.left_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
  if role=='caption':r.italic=True;p.paragraph_format.keep_with_next=True
  else:
   r.font.name='Courier New';r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0);r.italic=False;r.bold=False
   p.paragraph_format.line_spacing=1.05;p.paragraph_format.keep_together=True;p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(10)
 return p
anchor=next(p for p in d.paragraphs if p.style.name=='Heading 2' and p.text.startswith('4.5 '))
add(anchor,'4.4.5 Visual examples from the recorded experiments','head')
add(anchor,'Figures 4.1 and 4.2 illustrate two saved frames from the third completed OpenCV run, recorded on 9 September 2026 at 10:57:12. Each figure places the original camera image beside the white and yellow masks and the control overlay. The masks were reconstructed offline from the saved image using the archived threshold and morphology settings; the camera images and control overlays are the original experiment outputs. The grey area in each mask panel denotes the region outside processing, rather than rejected road pixels.')
for num,title,caption,explain in [
(30,'4.1','Straight-section example from OpenCV run 3, frame 000030.','In Figure 4.1, the yellow centre marking and right white boundary provide paired evidence. The logged steering request is 373, close to the calibrated centre of 365, while the forward request is 400. Red points in the saved overlay mark the centre observations passed to the controller; the green circle marks the near target and the magenta circle marks the farther target. These plotted points are algorithm outputs rather than manually labelled ground truth.'),
(104,'4.2','Left-bend example from OpenCV run 3, frame 000104.','Figure 4.2 shows a different lane shape as the vehicle enters the left bend. The relative positions of the near and far targets affect the preview term, and the logged steering request is 394 with the same forward request of 400. Replaying the archived detector on both selected raw frames reproduced their logged steering counts. The two examples explain the processing chain; they do not by themselves measure detection accuracy or establish performance throughout a lap.')]:
 p=add(anchor,'');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(8)
 p.add_run().add_picture(str(base/f'figure_{num:06d}.png'),width=Inches(6.1))
 c=add(anchor,f'Figure {title} {caption}','caption');c.paragraph_format.keep_with_next=False
 add(anchor,explain)
anchor=next(p for p in d.paragraphs if p.style.name=='Heading 1' and p.text.startswith('5.'))
add(anchor,'4.6.1 Key control code from the experimental runner','head')
add(anchor,'The following excerpts connect the image-space controller in Section 4.4 to the actuator requests used in the experiments. They are taken from the code snapshots saved with OpenCV run 3. Line wrapping is adjusted for presentation. The excerpts depend on their surrounding classes, imports and validity checks, and are not standalone vehicle-start commands.')
add(anchor,'Listing 4.1 Preview error and steering count from control() in s2_preview_core_local.py.','caption')
add(anchor,'''near_offset = (near - width / 2) / (width / 2)
result['near_offset'] = near_offset
for y in (0.55, 0.6):
    far, far_status = target(points, y)
    if far is not None:
        break
if far is None:
    return result
far_offset = (far - width / 2) / (width / 2)
trend = (far_offset - near_offset) * (0.1 / (0.65 - y))
preview = float(np.clip(near_offset + 1.5 * trend, -1, 1))
command = int(round(
    365 - 75 * float(np.clip(preview / 0.5, -1, 1))))''','code')
add(anchor,'The preceding part of control() requires a selected track and an available near target. Listing 4.1 then searches for the far target, computes the preview error and converts it to a bounded steering count. Returning without a far target leaves the primary branch unavailable; the detector may subsequently attempt its separately constrained estimation branches. This distinction prevents the primary formula from being applied to missing image evidence.')
add(anchor,'Listing 4.2 Start and stop gate from Gate.tick() in bounded_state_local_16s.py.','caption')
add(anchor,'''def tick(self, now, ready, fresh=True):
    if self.reason:
        return 370
    if self.started is not None:
        if now - self.started >= DURATION_SECONDS:
            self.stop('time_limit')
        elif not fresh:
            self.stop('stale_frame')
        elif not ready:
            self.stop('target_unavailable')
    else:
        self.streak = self.streak + 1 if ready and fresh else 0
        if self.streak >= 5:
            self.started = now
    return (400 if self.started is not None
            and self.reason is None else 370)''','code')
add(anchor,'Gate.stop() preserves the first stop reason. After a reason is latched, later valid frames do not automatically restart motion. Before motion begins, an invalid or stale frame resets the five-frame streak. DURATION_SECONDS is 16.0 in this runner. Freshness is calculated upstream from sensor timestamp progression, frame age and the capture-to-processing interval; this gate consumes that Boolean result.')
add(anchor,'Listing 4.3 Guarded output update from Output.update() in v2_local_continuous_16s.py.','caption')
add(anchor,'''def update(self, ready, steering, fresh):
    with self.lock:
        now = time.monotonic()
        if (self.gate.started is not None
                and now - self.last > .25):
            self.gate.stop('watchdog')
        self.last = now
        throttle = self.gate.tick(now, ready, fresh)
        steering = (steering if ready and fresh
                    and not self.gate.reason else 365)
        self.write(throttle, steering)
        return throttle, steering, self.gate.reason''','code')
add(anchor,'The write() method sends throttle to PCA9685 channel 0 and steering to channel 1 using set_pwm(channel, 0, count). When a stop has been latched, the gate supplies throttle 370 and this update selects steering centre 365. A separate watchdog thread also requests these neutral values when the main loop stops updating. These are software requests: an interrupted I2C connection can prevent delivery, which is why the experimental procedure retained physical disconnection of the motor supply as the fallback.')
# Insert two new TOC3 entries in heading order.
template=deepcopy(next(p._p for p in d.paragraphs if p.style.name.lower()=='toc 3'))
for i,h in enumerate([p for p in d.paragraphs if p.style.name=='Heading 3' and p.text.startswith(('4.4.5 ','4.6.1 '))]):
 anchor_name='VisualCodeSub'+str(i);bid=str(11000+i)
 b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),bid);b.set(qn('w:name'),anchor_name);h._p.insert(1,b)
 e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),bid);h._p.append(e)
 prefix='.'.join(h.text.split()[0].split('.')[:2])+' '
 parent=next(p for p in d.paragraphs if p.style.name.lower().startswith('toc') and p.text.startswith(prefix));cur=parent._p
 while cur.getnext() is not None:
  nxt=cur.getnext();s=nxt.find('./'+qn('w:pPr')+'/'+qn('w:pStyle'))
  if s is None or s.get(qn('w:val'))!='TOC3':break
  cur=nxt
 new=deepcopy(template);ts=new.findall('.//'+qn('w:t'));ts[0].text=h.text;ts[-1].text='1'
 for link in new.findall('.//'+qn('w:hyperlink')):link.set(qn('w:anchor'),anchor_name)
 for ins in new.findall('.//'+qn('w:instrText')):ins.text=' PAGEREF '+anchor_name+' \\h '
 cur.addnext(new)
d.save(out);print(out)
